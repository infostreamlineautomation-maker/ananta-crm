import io
import re
import sys
from datetime import datetime
from difflib import SequenceMatcher
import openpyxl
from pypdf import PdfReader


def clean_text(text: str) -> str:
    if not text:
        return ""
    return re.sub(r"[ \t]+", " ", text).strip()


def normalize_for_match(text: str) -> str:
    if not text:
        return ""
    t = text.lower()
    for noise in ["pvt", "ltd", "limited", "private", "corp", "inc", "co", "llp", "company"]:
        t = re.sub(rf"\b{noise}\.?\b", "", t)
    return re.sub(r"[^a-z0-9]", "", t)


def parse_date_str(date_str: str) -> str:
    if not date_str:
        return datetime.today().strftime("%Y-%m-%d")
    s = date_str.strip()
    patterns = [
        ("%d/%m/%Y", r"\d{1,2}/\d{1,2}/\d{4}"),
        ("%d-%m-%Y", r"\d{1,2}-\d{1,2}-\d{4}"),
        ("%d.%m.%Y", r"\d{1,2}\.\d{1,2}\.\d{4}"),
        ("%Y-%m-%d", r"\d{4}-\d{1,2}-\d{1,2}"),
        ("%d/%m/%y", r"\d{1,2}/\d{1,2}/\d{2}"),
        ("%d-%m-%y", r"\d{1,2}-\d{1,2}-\d{2}"),
    ]
    for fmt, regex in patterns:
        m = re.search(regex, s)
        if m:
            try:
                dt = datetime.strptime(m.group(0), fmt)
                if dt.year < 100:
                    dt = dt.replace(year=dt.year + 2000)
                return dt.strftime("%Y-%m-%d")
            except ValueError:
                continue
    return datetime.today().strftime("%Y-%m-%d")


def find_matching_client(client_name_query: str, organization=None):
    from apps.clients.models import Client

    if not client_name_query:
        return None

    qs = Client.objects.filter(is_deleted=False).select_related("company")
    if organization:
        qs = qs.filter(organization=organization)

    norm_query = normalize_for_match(client_name_query)
    if not norm_query:
        return None

    best_match = None
    best_score = 0.0

    for c in qs:
        comp_name = c.company.company_name if c.company else ""
        for candidate in [c.client_name, comp_name]:
            if not candidate:
                continue
            norm_cand = normalize_for_match(candidate)
            if not norm_cand:
                continue

            if norm_query == norm_cand:
                return {"id": c.id, "client_name": c.client_name, "company_name": comp_name}

            # Substring match
            if norm_query in norm_cand or norm_cand in norm_query:
                score = min(len(norm_query), len(norm_cand)) / max(len(norm_query), len(norm_cand))
                if score > best_score:
                    best_score = score
                    best_match = {"id": c.id, "client_name": c.client_name, "company_name": comp_name}

            # Similarity match
            ratio = SequenceMatcher(None, norm_query, norm_cand).ratio()
            if ratio > best_score and ratio >= 0.75:
                best_score = ratio
                best_match = {"id": c.id, "client_name": c.client_name, "company_name": comp_name}

    return best_match if best_score >= 0.70 else None


def find_matching_supplier(supplier_name_query: str, organization=None):
    from apps.suppliers.models import Supplier

    if not supplier_name_query:
        return None

    qs = Supplier.objects.filter(is_deleted=False)
    if organization:
        qs = qs.filter(organization=organization)

    norm_query = normalize_for_match(supplier_name_query)
    if not norm_query:
        return None

    best_match = None
    best_score = 0.0

    for s in qs:
        for candidate in [s.supplier_name, s.company_name]:
            if not candidate:
                continue
            norm_cand = normalize_for_match(candidate)
            if not norm_cand:
                continue

            if norm_query == norm_cand:
                return {"id": s.id, "supplier_name": s.supplier_name}

            if norm_query in norm_cand or norm_cand in norm_query:
                score = min(len(norm_query), len(norm_cand)) / max(len(norm_query), len(norm_cand))
                if score > best_score:
                    best_score = score
                    best_match = {"id": s.id, "supplier_name": s.supplier_name}

            ratio = SequenceMatcher(None, norm_query, norm_cand).ratio()
            if ratio > best_score and ratio >= 0.75:
                best_score = ratio
                best_match = {"id": s.id, "supplier_name": s.supplier_name}

    return best_match if best_score >= 0.70 else None


def is_header_line(text: str) -> bool:
    clean = re.sub(r"[^A-Za-z0-9]", " ", text).strip().upper()
    tokens = set(clean.split())
    header_keywords = {
        "SR", "NO", "DESCRIPTION", "PARTICULARS", "ITEM", "QTY", "QUANTITY",
        "RATE", "PRICE", "PIECE", "AMOUNT", "TOTAL", "PER"
    }
    if tokens and tokens.issubset(header_keywords):
        return True
    if len(tokens) <= 3 and any(k in tokens for k in ["SR", "NO", "DESCRIPTION", "QTY", "RATE", "PIECE", "AMOUNT"]):
        return True
    return False


def clean_number_str(s: str) -> str:
    if not s:
        return "0"
    cleaned = re.sub(r"[^\d.]", "", s)
    try:
        val = float(cleaned)
        return str(int(val)) if val.is_integer() else str(val)
    except ValueError:
        return cleaned or "0"


class PDFQuotationExtractor:
    def __init__(self, file_bytes: bytes, organization=None):
        self.file_bytes = file_bytes
        self.organization = organization

    def _extract_ocr_lines(self) -> list[str]:
        """
        Renders PDF pages to high-resolution images and uses OCR to extract
        text from scanned documents or PDFs where fonts were converted to vector paths.
        Supports winocr on Windows and pytesseract on Linux/any platform.
        """
        reconstructed_lines = []
        try:
            import pypdfium2 as pdfium
            pdf = pdfium.PdfDocument(io.BytesIO(self.file_bytes))
        except Exception:
            return []

        # 1. On Windows, try native Windows Media OCR if winocr is installed
        if sys.platform == "win32":
            try:
                import winocr

                for page in pdf:
                    pil_img = page.render(scale=2).to_pil()
                    result = winocr.recognize_pil_sync(pil_img, lang="en")
                    if not result or "lines" not in result:
                        continue

                    all_words = []
                    for line in result["lines"]:
                        for w in line.get("words", []):
                            rect = w.get("bounding_rect", {})
                            all_words.append({
                                "text": w["text"],
                                "x": rect.get("x", 0),
                                "y": rect.get("y", 0),
                                "w": rect.get("width", 0),
                                "h": rect.get("height", 0),
                            })

                    all_words.sort(key=lambda w: (w["y"], w["x"]))
                    rows = []
                    for w in all_words:
                        placed = False
                        for r in rows:
                            avg_y = sum(item["y"] for item in r) / len(r)
                            if abs(w["y"] - avg_y) <= 12:
                                r.append(w)
                                placed = True
                                break
                        if not placed:
                            rows.append([w])

                    rows.sort(key=lambda r: sum(item["y"] for item in r) / len(r))
                    for r in rows:
                        r.sort(key=lambda item: item["x"])
                        line_str = " ".join(item["text"] for item in r).strip()
                        if line_str:
                            reconstructed_lines.append(line_str)

                if reconstructed_lines:
                    return reconstructed_lines
            except Exception:
                pass

        # 2. Cross-platform / Linux fallback using pytesseract if available
        try:
            import pytesseract

            for page in pdf:
                pil_img = page.render(scale=2).to_pil()
                text = pytesseract.image_to_string(pil_img)
                if text:
                    for l in text.splitlines():
                        cl = l.strip()
                        if cl:
                            reconstructed_lines.append(cl)
            if reconstructed_lines:
                return reconstructed_lines
        except Exception:
            pass

        return reconstructed_lines

    def extract(self) -> dict:
        full_text = ""
        # 1. Primary text extraction using pypdfium2 (Google PDFium spatial engine)
        try:
            import pypdfium2 as pdfium
            doc = pdfium.PdfDocument(io.BytesIO(self.file_bytes))
            for page in doc:
                textpage = page.get_textpage()
                full_text += textpage.get_text_range() + "\n"
        except Exception:
            # Fallback to pypdf PdfReader
            try:
                reader = PdfReader(io.BytesIO(self.file_bytes))
                for page in reader.pages:
                    full_text += (page.extract_text() or "") + "\n"
            except Exception:
                pass

        lines = [line.strip() for line in full_text.splitlines() if line.strip()]

        # If PDF has vector-outlined text, raster images, or scanned graphics with no embedded fonts:
        if len(full_text.strip()) < 50:
            ocr_lines = self._extract_ocr_lines()
            if ocr_lines:
                lines = ocr_lines
                full_text = "\n".join(ocr_lines)

        # 1. Date extraction
        quotation_date = datetime.today().strftime("%Y-%m-%d")
        for line in lines:
            m = re.search(r"(?:Date|DATE)\s*[:\s\-]\s*(\d{1,2}[\/\-\.]\d{1,2}[\/\-\.]\d{2,4})", line)
            if m:
                quotation_date = parse_date_str(m.group(1))
                break

        # 2. Recipient extraction (To, ...)
        to_name = ""
        to_address_lines = []
        to_idx = -1
        for i, line in enumerate(lines):
            if re.match(r"^To\b", line, re.IGNORECASE):
                to_idx = i
                break

        if to_idx != -1 and to_idx + 1 < len(lines):
            raw_to = lines[to_idx + 1]
            to_name = re.sub(r"(?:Date|DATE|Quotation\s*No\.?).*$", "", raw_to, flags=re.IGNORECASE).strip()
            for j in range(to_idx + 2, min(to_idx + 6, len(lines))):
                l = lines[j]
                if re.match(r"^(SUB|Sub|SUBJECT|Subject|Date|DATE|With The Above)\b", l):
                    break
                to_address_lines.append(l)

        to_address = "\n".join(to_address_lines)

        # 3. Subject extraction
        subject = ""
        for line in lines:
            m = re.search(r"^(?:SUB|Sub|SUBJECT|Subject)\.?\s*[:\s\-]\s*(.*)$", line)
            if m:
                raw_subj = m.group(1).strip()
                subject = re.sub(r"^[.\s:]+", "", raw_subj)
                if subject.endswith("."):
                    subject = subject[:-1].strip()
                break

        # 4. Intro text extraction
        intro_text = ""
        for i, line in enumerate(lines):
            if "with the above reference" in line.lower():
                intro_parts = [line]
                for next_line in lines[i + 1 : i + 5]:
                    if re.search(r"\b(SR\.?\s*NO|DESCRIPTION|QTY|RATE)\b", next_line, re.IGNORECASE):
                        break
                    intro_parts.append(next_line)
                intro_text = " ".join(intro_parts)
                intro_text = re.sub(r"\s+SR\.?$", "", intro_text).strip()
                break

        # 5. Table extraction
        items = []
        table_start_idx = -1
        for i, line in enumerate(lines):
            has_table_cols = re.search(r"\b(DESCRIPTION|PARTICULARS|ITEM)\b", line, re.IGNORECASE) and re.search(r"\b(QTY|RATE|PRICE|AMOUNT)\b", line, re.IGNORECASE)
            has_sr_header = bool(re.search(r"\bSR\.?\s*NO\b", line, re.IGNORECASE) or ("SR." in line and i + 1 < len(lines) and "NO." in lines[i + 1]))
            if has_table_cols or (has_sr_header and re.search(r"\b(QTY|RATE|DESCRIPTION)\b", line, re.IGNORECASE)):
                table_start_idx = i
                break

        row_pat = re.compile(
            r"^(?:(?P<sr>\d+)[\.\s]+)?(?P<desc>.+?)\s+(?P<qty>[0-9,]+(?:\.[0-9]+)?)\s+(?P<rate>[0-9,]+(?:\.[0-9]+)?)(?:\s+(?P<amount>.*))?$"
        )
        qty_rate_pat = re.compile(
            r"^(?:(?P<desc_cont>.*?)\s+)?(?P<qty>[0-9,]+(?:\.[0-9]+)?)\s+(?P<rate>[0-9,]+(?:\.[0-9]+)?)(?:\s+(?P<amount>.*))?$"
        )

        if table_start_idx != -1:
            pending_desc = ""
            for j in range(table_start_idx, len(lines)):
                l = lines[j]
                if is_header_line(l):
                    continue
                # Stop if reaching notes or footer or subtotal
                if re.match(r"^(NOTES|Notes|TERMS|Terms|Thank you|Thanking you|Total|TOTAL|SUBTOTAL|Subtotal)\b", l, re.IGNORECASE):
                    break

                row_match = row_pat.match(l)
                if row_match:
                    desc = (pending_desc + " " + row_match.group("desc")).strip()
                    qty_str = clean_number_str(row_match.group("qty"))
                    rate_str = clean_number_str(row_match.group("rate"))
                    amt_val = round(float(qty_str) * float(rate_str), 2)
                    items.append({
                        "description": desc,
                        "qty": qty_str,
                        "rate": rate_str,
                        "supplier_rate": "0",
                        "margin_percent": "25",
                        "profit": "0",
                        "total_profit": "0",
                        "amount": str(amt_val),
                        "extra_data": {},
                    })
                    pending_desc = ""
                else:
                    qm = qty_rate_pat.match(l)
                    if qm and pending_desc:
                        cont = qm.group("desc_cont") or ""
                        desc = (pending_desc + " " + cont).strip()
                        qty_str = clean_number_str(qm.group("qty"))
                        rate_str = clean_number_str(qm.group("rate"))
                        amt_val = round(float(qty_str) * float(rate_str), 2)
                        items.append({
                            "description": desc,
                            "qty": qty_str,
                            "rate": rate_str,
                            "supplier_rate": "0",
                            "margin_percent": "25",
                            "profit": "0",
                            "total_profit": "0",
                            "amount": str(amt_val),
                            "extra_data": {},
                        })
                        pending_desc = ""
                    else:
                        cleaned_l = re.sub(r"^\d+[\.\s]+", "", l).strip()
                        if cleaned_l and not any(cleaned_l.lower().startswith(p) for p in ["subtotal", "total", "notes", "terms", "gst"]):
                            pending_desc = (pending_desc + " " + cleaned_l).strip()

        # 6. Notes extraction
        notes = ""
        notes_match = re.search(
            r"(?:^|\n)\s*(?:NOTES|Notes|TERMS|Terms)\s*[:\-]\s*([\s\S]*?)(?:Thank you|Thanking you|Sincerely|For |Ananta|$)",
            full_text,
        )
        if notes_match:
            raw_notes = notes_match.group(1).strip()
            clean_lines = [l.strip() for l in raw_notes.splitlines() if l.strip()]
            notes = "\n".join(clean_lines)
        else:
            notes_start = -1
            for i, l in enumerate(lines):
                if re.match(r"^(?:NOTES|Notes|TERMS|Terms)\s*[:\-]?", l):
                    notes_start = i
                    break
            if notes_start != -1:
                note_lines = []
                for l in lines[notes_start + 1 :]:
                    if re.match(r"^(?:Thank you|Thanking you|Sincerely|For |Ananta)\b", l, re.IGNORECASE):
                        break
                    note_lines.append(l)
                notes = "\n".join(note_lines)

        # 7. Matching CRM client
        matched_client = find_matching_client(to_name, self.organization)

        # Compute missing fields
        missing_fields = []
        if not matched_client:
            missing_fields.append({
                "field": "client",
                "label": "Client Selection",
                "reason": f"Recipient '{to_name}' not automatically matched in CRM database",
            })
        missing_fields.append({
            "field": "supplier",
            "label": "Supplier Name",
            "reason": "Not specified in customer PDF quotation",
        })
        missing_fields.append({
            "field": "supplier_rate",
            "label": "Supplier Cost Rate",
            "reason": "PDF only specifies Customer Selling Rate. Enter Supplier Cost or Margin % to calculate profit.",
        })

        return {
            "source_type": "pdf",
            "document_profile": "customer_quotation",
            "data": {
                "quotation_date": quotation_date,
                "to_name": to_name,
                "to_address": to_address,
                "subject": subject,
                "intro_text": intro_text,
                "notes": notes,
                "client_id": matched_client["id"] if matched_client else None,
                "client_name": matched_client["client_name"] if matched_client else to_name,
                "supplier_id": None,
                "items": items,
            },
            "missing_fields": missing_fields,
        }


class ExcelQuotationExtractor:
    def __init__(self, file_bytes: bytes, organization=None):
        self.file_bytes = file_bytes
        self.organization = organization
        self.wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        self.ws = self.wb.active

    def extract(self) -> dict:
        sheet = self.ws
        all_rows = list(sheet.iter_rows(values_only=True))

        # Check for Costing Breakdown Profile (like Broucher)
        has_product_cost = False
        product_cost_val = 0.0
        total_cost_val = 0.0
        total_qty_val = 0.0
        costing_subitems = []
        job_title = ""

        # Scan for job title in first 5 rows
        for r_idx in range(min(5, len(all_rows))):
            row = all_rows[r_idx]
            for val in row:
                if val and isinstance(val, str) and len(val.strip()) > 2:
                    if not any(header in val.lower() for header in ["table", "column", "sr", "no", "qty", "rate"]):
                        job_title = val.strip()
                        break
            if job_title:
                break

        # Check for summary labels in table
        for r_idx, row in enumerate(all_rows):
            row_str = " ".join([str(v) for v in row if v is not None]).lower()
            if "product cost" in row_str:
                has_product_cost = True
                nums = [v for v in row if isinstance(v, (int, float))]
                if nums:
                    product_cost_val = float(nums[-1])
            if "total rs" in row_str or "total cost" in row_str:
                nums = [v for v in row if isinstance(v, (int, float))]
                if nums:
                    total_cost_val = float(nums[-1])
            if "total qty" in row_str or "total quantity" in row_str:
                nums = [v for v in row if isinstance(v, (int, float))]
                if nums:
                    total_qty_val = float(nums[-1])

        # If it's a Costing Breakdown Sheet (Broucher format):
        if has_product_cost or ("total rs" in str(all_rows).lower() and total_qty_val > 0):
            for row in all_rows:
                if not row or len(row) < 3:
                    continue
                name_cand = row[0] or (row[1] if len(row) > 1 else None)
                if name_cand and isinstance(name_cand, str):
                    name_clean = name_cand.strip()
                    if name_clean.lower() in ["column 1", "sr. no", "particulars", "description", "item", "total", "total rs.", "total qty.", "product cost"]:
                        continue
                    num_vals = [v for v in row if isinstance(v, (int, float)) and v > 0]
                    if len(num_vals) >= 2:
                        costing_subitems.append({
                            "name": name_clean,
                            "qty": num_vals[0],
                            "rate": num_vals[1] if len(num_vals) > 1 else 0,
                            "total": num_vals[-1] if len(num_vals) > 2 else (num_vals[0] * num_vals[1]),
                        })

            qty_final = total_qty_val if total_qty_val > 0 else 1000
            supplier_rate_final = product_cost_val if product_cost_val > 0 else (round(total_cost_val / qty_final, 2) if qty_final > 0 else 0)

            # Auto-calculate client rate with default 25% margin
            margin_percent = 25.0
            profit_unit = round(supplier_rate_final * (margin_percent / 100.0), 2)
            client_rate = round(supplier_rate_final + profit_unit, 2)
            total_profit = round(profit_unit * qty_final, 2)

            subitem_summary = ", ".join([f"{item['name']} ({item['qty']} @ ₹{item['rate']})" for item in costing_subitems])

            item_obj = {
                "description": job_title or "Custom Job Order",
                "qty": str(int(qty_final)) if float(qty_final).is_integer() else str(qty_final),
                "supplier_rate": str(supplier_rate_final),
                "margin_percent": str(margin_percent),
                "profit": str(profit_unit),
                "rate": str(client_rate),
                "total_profit": str(total_profit),
                "amount": str(round(client_rate * qty_final, 2)),
                "extra_data": {"specifications": subitem_summary} if subitem_summary else {},
            }

            return {
                "source_type": "excel",
                "document_profile": "costing_breakdown",
                "data": {
                    "quotation_date": datetime.today().strftime("%Y-%m-%d"),
                    "to_name": "",
                    "to_address": "",
                    "subject": f"QUOTATION FOR {job_title.upper()}" if job_title else "QUOTATION FOR PRINTING",
                    "intro_text": "With The Above Reference, We Are Pleased To Supply The Material With The Following Rates As Per Terms & Conditions Stated Below.",
                    "notes": "1. GST Should be Extra as per Government Rules.\n2. Transportation charge will be extra as per the Location of Delivery.\n3. Payment to be made by Payees A/c. Cheque, Draft or NEFT / RTGS / IMPS only.\n4. Subject to Surat Jurisdiction.",
                    "client_id": None,
                    "client_name": "",
                    "supplier_id": None,
                    "items": [item_obj],
                },
                "missing_fields": [
                    {"field": "client", "label": "Client Selection", "reason": "No client specified in internal costing sheet"},
                    {"field": "supplier", "label": "Supplier Name", "reason": "Select fulfilling supplier for this job"},
                    {"field": "margin_percent", "label": "Profit Margin %", "reason": f"Defaulted to 25% on unit cost ₹{supplier_rate_final}. Adjust if needed."},
                ],
            }

        # Case 2: Multi-Item Quotation / Supplier Rate Table (e.g. Quotationsample.xlsx)
        def map_headers(rows):
            best_map = {}
            best_last_header_idx = -1
            best_score = 0

            for r_idx in range(min(12, len(rows))):
                row = rows[r_idx]
                # Test single row
                cur_map = {}
                for c_idx, val in enumerate(row):
                    h = str(val).lower().strip() if val is not None else ""
                    if not h:
                        continue
                    if any(k in h for k in ["product", "description", "item", "particulars"]) and not any(k in h for k in ["qty", "rate", "cost"]):
                        cur_map["description"] = c_idx
                    elif any(k in h for k in ["supplier rate", "pur rate", "cost rate"]):
                        cur_map["supplier_rate"] = c_idx
                    elif ("client" in h and "rate" in h) or ("sell" in h and "rate" in h) or (h in ["rate", "rate/pc", "unit rate"]):
                        cur_map["rate"] = c_idx
                    elif any(k in h for k in ["qty", "quantity"]) and "rate" not in h:
                        cur_map["qty"] = c_idx
                    elif "total profit" in h:
                        cur_map["total_profit"] = c_idx
                    elif "profit" in h:
                        cur_map["profit"] = c_idx
                    elif any(k in h for k in ["%", "margin", "add as per %"]):
                        cur_map["margin_percent"] = c_idx

                score = ("description" in cur_map) + ("qty" in cur_map) + ("rate" in cur_map) + ("supplier_rate" in cur_map)
                if score > best_score and "qty" in cur_map:
                    best_score = score
                    best_map = cur_map
                    best_last_header_idx = r_idx

                # Test two-row header (r_idx and r_idx + 1)
                if r_idx + 1 < len(rows):
                    r_next = rows[r_idx + 1]
                    filled_r = []
                    curr = ""
                    for val in row:
                        if val is not None and str(val).strip():
                            curr = str(val).strip()
                        filled_r.append(curr)

                    two_row_map = {}
                    for c_idx, (p, sub) in enumerate(zip(filled_r, r_next)):
                        sub_str = str(sub).strip().lower() if sub is not None else ""
                        p_str = p.strip().lower() if p else ""
                        combined = f"{p_str} {sub_str}".strip()
                        if not combined:
                            continue

                        if any(k in sub_str for k in ["product", "description", "item", "particulars"]) or (
                            any(k in p_str for k in ["product", "description", "item", "particulars"])
                            and not any(k in sub_str for k in ["qty", "rate", "cost", "profit", "%"])
                        ):
                            two_row_map["description"] = c_idx
                        elif any(k in sub_str for k in ["qty", "quantity"]) or (
                            any(k in combined for k in ["qty", "quantity"]) and "rate" not in sub_str
                        ):
                            two_row_map["qty"] = c_idx
                        elif "total profit" in combined or "total profit" in sub_str:
                            two_row_map["total_profit"] = c_idx
                        elif "profit" in sub_str or ("profit" in combined and "total" not in combined):
                            two_row_map["profit"] = c_idx
                        elif any(k in sub_str for k in ["%", "margin", "add as per %"]) or any(k in combined for k in ["%", "margin", "add as per %"]):
                            two_row_map["margin_percent"] = c_idx
                        elif any(k in sub_str for k in ["rate", "price"]):
                            if "supplier" in p_str or "cost" in p_str or "pur" in p_str or "supplier" in sub_str:
                                two_row_map["supplier_rate"] = c_idx
                            else:
                                two_row_map["rate"] = c_idx
                        elif "rate" in combined and "supplier" not in combined and "profit" not in combined:
                            two_row_map["rate"] = c_idx

                    two_score = (
                        ("description" in two_row_map)
                        + ("qty" in two_row_map)
                        + ("rate" in two_row_map)
                        + ("supplier_rate" in two_row_map)
                        + ("profit" in two_row_map)
                        + ("margin_percent" in two_row_map)
                    )
                    if two_score > best_score and "qty" in two_row_map:
                        best_score = two_score
                        best_map = two_row_map
                        best_last_header_idx = r_idx + 1

            return best_map, best_last_header_idx

        col_map, header_row_idx = map_headers(all_rows)

        # Check top rows for Client / Supplier / Subject
        extracted_client_name = ""
        extracted_supplier_name = ""
        detected_subject = ""
        for r_idx in range(header_row_idx if header_row_idx != -1 else 10):
            row_str = " ".join([str(v) for v in all_rows[r_idx] if v is not None])
            m_client = re.search(r"(?:Client(?:\s*Name)?|Customer|To)\s*[:\-]\s*([A-Za-z0-9\s\.\&]+)", row_str, re.IGNORECASE)
            if m_client and not extracted_client_name:
                extracted_client_name = clean_text(m_client.group(1))

            m_supp = re.search(r"Supplier(?:\s*Name)?\s*[:\-]\s*([A-Za-z0-9\s\.\&]+)", row_str, re.IGNORECASE)
            if m_supp and not extracted_supplier_name:
                extracted_supplier_name = clean_text(m_supp.group(1))

        # Parse item rows
        items = []
        last_desc = ""
        proposal_lines = []

        if header_row_idx != -1:
            for r_idx in range(header_row_idx + 1, len(all_rows)):
                row = all_rows[r_idx]
                if not row or not any(row):
                    continue

                row_str = " ".join([str(v).strip() for v in row if v is not None]).strip()
                if not detected_subject:
                    m_qf = re.search(r"Quotation\s+For\s*[:\-]\s*([^\n\r]+)", row_str, re.IGNORECASE)
                    if m_qf:
                        detected_subject = "QUOTATION FOR " + m_qf.group(1).strip().upper()

                if any(row_str.lower().startswith(prefix) for prefix in ["dear", "surat", "quotation for", "thank you", "with reference", "box printing", "rs."]):
                    proposal_lines.append(row_str)
                    continue

                desc_idx = col_map.get("description", 1 if len(row) > 1 else 0)
                qty_idx = col_map.get("qty", 2 if len(row) > 2 else 1)
                rate_idx = col_map.get("rate")
                supp_rate_idx = col_map.get("supplier_rate")
                margin_idx = col_map.get("margin_percent")
                profit_idx = col_map.get("profit")
                size_idx = col_map.get("size")

                raw_desc = str(row[desc_idx]).strip() if desc_idx < len(row) and row[desc_idx] is not None else ""
                raw_desc = re.sub(r"[\xa0\s]+", " ", raw_desc).strip()
                if raw_desc.lower() in ["total", "subtotal", "notes", "terms", "grand total", "total rs.", "total qty.", "sr. no.", "product"]:
                    continue

                desc = raw_desc or last_desc
                if not desc:
                    continue

                qty_raw = row[qty_idx] if qty_idx < len(row) and row[qty_idx] is not None else 0
                try:
                    qty_num = float(qty_raw)
                    if qty_num <= 0:
                        continue
                    qty_str = str(int(qty_num)) if qty_num.is_integer() else str(qty_num)
                except (ValueError, TypeError):
                    continue

                last_desc = desc

                supp_rate = "0"
                if supp_rate_idx is not None and supp_rate_idx < len(row) and row[supp_rate_idx] is not None:
                    try:
                        supp_rate = str(float(row[supp_rate_idx]))
                    except (ValueError, TypeError):
                        pass

                rate = "0"
                if rate_idx is not None and rate_idx < len(row) and row[rate_idx] is not None:
                    try:
                        rate = str(float(row[rate_idx]))
                    except (ValueError, TypeError):
                        pass

                margin = "25"
                if margin_idx is not None and margin_idx < len(row) and row[margin_idx] is not None:
                    try:
                        m_val = float(row[margin_idx])
                        if m_val < 1.0 and m_val > 0:
                            m_val = m_val * 100.0
                        margin = str(round(m_val, 1))
                    except (ValueError, TypeError):
                        pass

                # Calculate profit / rate reactive math
                s_float = float(supp_rate)
                r_float = float(rate)
                m_float = float(margin)

                if s_float > 0 and r_float == 0:
                    profit_unit = round(s_float * (m_float / 100.0), 2)
                    r_float = round(s_float + profit_unit, 2)
                    rate = str(r_float)
                elif r_float > 0 and s_float == 0:
                    s_float = round(r_float / (1.0 + (m_float / 100.0)), 2)
                    supp_rate = str(s_float)

                profit_unit = max(0.0, round(r_float - s_float, 2))
                tot_profit = round(profit_unit * qty_num, 2)

                extra_data = {}
                if size_idx is not None and size_idx < len(row) and row[size_idx] is not None:
                    extra_data["size"] = str(row[size_idx]).strip()

                items.append({
                    "description": desc,
                    "qty": qty_str,
                    "supplier_rate": supp_rate,
                    "margin_percent": margin,
                    "profit": str(profit_unit),
                    "rate": rate,
                    "total_profit": str(tot_profit),
                    "amount": str(round(qty_num * r_float, 2)),
                    "extra_data": extra_data,
                })

        # Match Client & Supplier
        matched_client = find_matching_client(extracted_client_name, self.organization) if extracted_client_name else None
        matched_supplier = find_matching_supplier(extracted_supplier_name, self.organization) if extracted_supplier_name else None

        missing_fields = []
        if not matched_client:
            missing_fields.append({
                "field": "client",
                "label": "Client Selection",
                "reason": f"Client '{extracted_client_name}' not matched in CRM" if extracted_client_name else "No client found in spreadsheet",
            })
        if not matched_supplier:
            missing_fields.append({
                "field": "supplier",
                "label": "Supplier Name",
                "reason": f"Supplier '{extracted_supplier_name}' not matched in CRM" if extracted_supplier_name else "No supplier found in spreadsheet",
            })

        first_desc = items[0]["description"].upper() if items else "PRINTING"
        final_subject = detected_subject or f"QUOTATION FOR {first_desc}"

        return {
            "source_type": "excel",
            "document_profile": "quotation_sheet",
            "data": {
                "quotation_date": datetime.today().strftime("%Y-%m-%d"),
                "to_name": extracted_client_name,
                "to_address": "",
                "subject": final_subject,
                "intro_text": "With The Above Reference, We Are Pleased To Supply The Material With The Following Rates As Per Terms & Conditions Stated Below.",
                "notes": "1. GST Should be Extra as per Government Rules.\n2. Transportation charge will be extra as per the Location of Delivery.\n3. Payment to be made by Payees A/c. Cheque, Draft or NEFT / RTGS / IMPS only.\n4. Subject to Surat Jurisdiction.",
                "client_id": matched_client["id"] if matched_client else None,
                "client_name": matched_client["client_name"] if matched_client else extracted_client_name,
                "supplier_id": matched_supplier["id"] if matched_supplier else None,
                "supplier_name": matched_supplier["supplier_name"] if matched_supplier else extracted_supplier_name,
                "items": items,
            },
            "missing_fields": missing_fields,
        }

